import sys
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = sys.argv[1]
MONTHS = 36
FIRST_COL = 3  # column C = month 1

BLUE = Font(name="Arial", size=10, color="0000FF")
BLACK = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
GREEN = Font(name="Arial", size=10, color="008000")
TITLE = Font(name="Arial", size=14, bold=True)
HEAD = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F7A55")
SECTION_FILL = PatternFill("solid", fgColor="DFF2E8")
YELLOW = PatternFill("solid", fgColor="FFFF00")
TOP = Border(top=Side(style="thin", color="999999"))

USD = '$#,##0;($#,##0);-'
NUM = '#,##0;(#,##0);-'
NUM1 = '#,##0.0;(#,##0.0);-'
PCT = '0.0%;(0.0%);-'

wb = Workbook()

# ---------------------------------------------------------------- Assumptions
a = wb.active
a.title = "Assumptions"
a.column_dimensions["A"].width = 46
a.column_dimensions["B"].width = 14
a.column_dimensions["C"].width = 14
a.column_dimensions["D"].width = 60
a["A1"] = "Closebook — Financial Model Assumptions"
a["A1"].font = TITLE
a["A2"] = "Legend: blue = input you can change · yellow fill = key assumption or placeholder to confirm · black = formula · green = link to another sheet. All figures are founder estimates, not sourced data."
a["A2"].font = Font(name="Arial", size=9, italic=True)
a["A2"].alignment = Alignment(wrap_text=True)
a.merge_cells("A2:D2")
a.row_dimensions[2].height = 30

ref = {}
row = 4


def section(title):
    global row
    row += 1
    for col in range(1, 5):
        a.cell(row=row, column=col).fill = SECTION_FILL
    a.cell(row=row, column=1, value=title).font = BOLD
    row += 1


def inp(key, label, value, fmt, note="", key_assumption=False):
    global row
    a.cell(row=row, column=1, value=label).font = BLACK
    c = a.cell(row=row, column=2, value=value)
    c.font = BLUE
    c.number_format = fmt
    if key_assumption:
        c.fill = YELLOW
    a.cell(row=row, column=4, value=note).font = Font(name="Arial", size=9, color="555555")
    ref[key] = f"Assumptions!$B${row}"
    row += 1


section("Pricing ($ per client per month)")
inp("p_starter", "Starter plan price", 299, USD, "From pitch deck pricing.")
inp("p_growth", "Growth plan price", 499, USD, "From pitch deck pricing.")
inp("p_pro", "Pro plan price", 899, USD, "From pitch deck pricing.")
inp("m_starter", "Plan mix: Starter", 0.50, PCT, "Share of clients on each plan; must sum to 100%.", True)
inp("m_growth", "Plan mix: Growth", 0.35, PCT, "", True)
inp("m_pro", "Plan mix: Pro", 0.15, PCT, "", True)
a.cell(row=row, column=1, value="Plan mix check (must be 100%)").font = BLACK
c = a.cell(row=row, column=2, value=f"={ref['m_starter'].split('!')[1]}+{ref['m_growth'].split('!')[1]}+{ref['m_pro'].split('!')[1]}")
c.number_format = PCT
c.font = BLACK
row += 1
a.cell(row=row, column=1, value="Blended price per client (ARPA)").font = BOLD
c = a.cell(
    row=row,
    column=2,
    value="=" + "+".join(f"{ref[p].split('!')[1]}*{ref[m].split('!')[1]}" for p, m in [("p_starter", "m_starter"), ("p_growth", "m_growth"), ("p_pro", "m_pro")]),
)
c.number_format = USD
c.font = BOLD
ref["arpa"] = f"Assumptions!$B${row}"
row += 1
inp("cu_attach", "Catch-up attach rate (share of new clients)", 0.40, PCT, "Clients who buy one-time catch-up bookkeeping when they join.", True)
inp("cu_price", "Average catch-up price", 2500, USD, "Deck range: $1,500–$5,000 one-time.")

section("Customer acquisition")
inp("d_start", "Direct new clients in month 1", 5, NUM, "Concierge phase: hand-served Shopify sellers.", True)
inp("g_y1", "Direct new-client growth per month, year 1", 0.15, PCT, "Month-over-month growth in new clients.", True)
inp("g_y2", "Direct new-client growth per month, year 2", 0.06, PCT, "", True)
inp("g_y3", "Direct new-client growth per month, year 3", 0.03, PCT, "", True)
inp("cac", "Direct CAC (marketing $ per new direct client)", 1500, USD, "Paid + content cost to win one direct client.", True)
inp("mkt_fixed", "Fixed marketing spend per month", 2000, USD, "Content, tools, events.")
inp("pt_start", "CPA partner program starts (month #)", 7, NUM, "Deck roadmap: white-label in months 9–18; pilots earlier.")
inp("pt_y1", "New CPA partner firms per month, year 1", 1, NUM1, "Applies from the start month.", True)
inp("pt_y2", "New CPA partner firms per month, year 2", 2, NUM1, "", True)
inp("pt_y3", "New CPA partner firms per month, year 3", 3, NUM1, "", True)
inp("pt_clients", "New clients per active partner per month", 1.0, NUM1, "Deck: one firm can bring 20–100 clients over time.", True)
inp("pt_share", "Revenue share paid to partner firms", 0.20, PCT, "Of subscription revenue from partner-sourced clients.")
inp("churn", "Monthly client churn", 0.03, PCT, "Deck target: under 2%; modelled conservatively.", True)

section("Cost of revenue")
inp("c_cpa", "CPA review cost per client per month", 20, USD, "~30 min/month at ~$40/hr (tech spec target).")
inp("c_ai", "AI / LLM cost per client per month", 25, USD, "Deck range: $15–30.")
inp("c_data", "Bank and data feed cost per client per month", 8, USD, "Deck range: $5–10.")
inp("cu_cost", "Catch-up delivery cost (% of catch-up revenue)", 0.35, PCT, "Contract bookkeeper time for back months.")
inp("pay_fee", "Payment processing (% of revenue)", 0.03, PCT, "Card processing on subscriptions.")

section("Operating expenses")
inp("burden", "Payroll burden (taxes and benefits, % of salary)", 0.20, PCT, "")
inp("ga_base", "G&A base per month (software, legal, insurance)", 4000, USD, "")
inp("ga_client", "G&A per client per month (support tools)", 5, USD, "")

section("Funding")
inp("cash0", "Starting cash", 50000, USD, "Founders' cash before the round. Placeholder.", True)
inp("raise", "Pre-seed raise amount", 1500000, USD, "Deck shows [$__]: replace with the real target.", True)
inp("raise_m", "Month the raise closes (month #)", 1, NUM, "", True)

section("Team (monthly salary, start month)")
a.cell(row=row, column=1, value="Role").font = BOLD
a.cell(row=row, column=2, value="Salary / mo").font = BOLD
a.cell(row=row, column=3, value="Start month #").font = BOLD
row += 1
hires = [
    ("Founder / CEO", 10000, 1),
    ("Founder / CTO", 10000, 1),
    ("Engineer #1", 13000, 4),
    ("Customer success + CPA lead", 9000, 6),
    ("Partnerships lead", 10000, 9),
    ("Engineer #2", 13000, 10),
    ("Marketing lead", 9000, 13),
    ("Engineer #3", 13000, 16),
    ("Customer success #2", 7500, 20),
    ("Customer success #3", 7500, 26),
    ("Finance + ops lead", 11000, 28),
    ("Engineer #4", 13000, 25),
]
h_first = row
for role, sal, start in hires:
    a.cell(row=row, column=1, value=role).font = BLUE
    c = a.cell(row=row, column=2, value=sal)
    c.font = BLUE
    c.number_format = USD
    c = a.cell(row=row, column=3, value=start)
    c.font = BLUE
    c.number_format = NUM
    row += 1
h_last = row - 1
a.cell(row=row, column=1, value="Add hires by inserting rows inside this table so the payroll formula picks them up.").font = Font(name="Arial", size=9, italic=True)
SAL = f"Assumptions!$B${h_first}:$B${h_last}"
START = f"Assumptions!$C${h_first}:$C${h_last}"

# ---------------------------------------------------------------- Monthly
m = wb.create_sheet("Monthly")
m.column_dimensions["A"].width = 40
m.column_dimensions["B"].width = 4
for i in range(MONTHS):
    m.column_dimensions[get_column_letter(FIRST_COL + i)].width = 11
m["A1"] = "Closebook — Monthly Projection (36 months, $)"
m["A1"].font = TITLE
m.freeze_panes = "C4"

rows = {}
r = 3


def L(i):
    return get_column_letter(FIRST_COL + i)


def line(key, label, formula, fmt=USD, bold=False, top=False):
    """formula(i, col) returns the cell formula for month index i (0-based)."""
    global r
    rows[key] = r  # set first so a row can refer to its own previous month
    m.cell(row=r, column=1, value=label).font = BOLD if bold else BLACK
    for i in range(MONTHS):
        c = m.cell(row=r, column=FIRST_COL + i, value=formula(i, L(i)))
        c.number_format = fmt
        uses_link = "Assumptions!" in str(c.value) and not any(ch in str(c.value) for ch in "+-*/(")
        c.font = GREEN if uses_link else (BOLD if bold else BLACK)
        if top:
            c.border = TOP
    if top:
        m.cell(row=r, column=1).border = TOP
    rows[key] = r
    r += 1


def head(title):
    global r
    r += 1
    m.cell(row=r, column=1, value=title).font = BOLD
    for col in range(1, FIRST_COL + MONTHS):
        m.cell(row=r, column=col).fill = SECTION_FILL
    r += 1


def span(k1, k2):
    return ref[k1] + ":" + ref[k2].split("!")[1]


def R(key, col):
    return f"{col}{rows[key]}"


def prev(key, i):
    return f"{L(i - 1)}{rows[key]}"


# header rows
m.cell(row=3, column=1, value="Month #").font = HEAD
m.cell(row=3, column=1).fill = HEAD_FILL
m.cell(row=3, column=2).fill = HEAD_FILL
for i in range(MONTHS):
    c = m.cell(row=3, column=FIRST_COL + i, value=i + 1)
    c.font = HEAD
    c.fill = HEAD_FILL
    c.alignment = Alignment(horizontal="center")
rows["month"] = 3
r = 4
line("year", "Year", lambda i, c: f"=ROUNDUP({c}$3/12,0)", '0')

head("Customers")
line("d_new", "New direct clients",
     lambda i, c: f"={ref['d_start']}" if i == 0 else f"={prev('d_new', i)}*(1+INDEX({span('g_y1','g_y3')},{c}$4))", NUM1)
line("partners", "Active CPA partner firms",
     lambda i, c: ("=0" if i == 0 else f"={prev('partners', i)}")
     + f"+IF({c}$3>={ref['pt_start']},INDEX({span('pt_y1','pt_y3')},{c}$4),0)", NUM1)
line("p_new", "New partner-sourced clients", lambda i, c: f"={R('partners', c)}*{ref['pt_clients']}", NUM1)
line("new", "Total new clients", lambda i, c: f"={R('d_new', c)}+{R('p_new', c)}", NUM1, bold=True)
line("d_end", "Direct clients (end of month)",
     lambda i, c: f"={R('d_new', c)}" if i == 0 else f"={prev('d_end', i)}*(1-{ref['churn']})+{R('d_new', c)}", NUM1)
line("p_end", "Partner clients (end of month)",
     lambda i, c: f"={R('p_new', c)}" if i == 0 else f"={prev('p_end', i)}*(1-{ref['churn']})+{R('p_new', c)}", NUM1)
line("clients", "Total clients (end of month)", lambda i, c: f"={R('d_end', c)}+{R('p_end', c)}", NUM1, bold=True, top=True)
line("churned", "Clients churned",
     lambda i, c: "=0" if i == 0 else f"=({prev('d_end', i)}+{prev('p_end', i)})*{ref['churn']}", NUM1)

head("Revenue")
line("subs", "Subscription revenue (MRR)", lambda i, c: f"={R('clients', c)}*{ref['arpa']}")
line("catchup", "Catch-up revenue (one-time)", lambda i, c: f"={R('new', c)}*{ref['cu_attach']}*{ref['cu_price']}")
line("share", "Less: partner revenue share", lambda i, c: f"=-{R('p_end', c)}*{ref['arpa']}*{ref['pt_share']}")
line("rev", "Net revenue", lambda i, c: f"={R('subs', c)}+{R('catchup', c)}+{R('share', c)}", bold=True, top=True)

head("Cost of revenue")
line("c_review", "CPA review", lambda i, c: f"={R('clients', c)}*{ref['c_cpa']}")
line("c_ai", "AI / LLM", lambda i, c: f"={R('clients', c)}*{ref['c_ai']}")
line("c_data", "Bank and data feeds", lambda i, c: f"={R('clients', c)}*{ref['c_data']}")
line("c_cu", "Catch-up delivery", lambda i, c: f"={R('catchup', c)}*{ref['cu_cost']}")
line("c_pay", "Payment processing", lambda i, c: f"={R('rev', c)}*{ref['pay_fee']}")
line("cogs", "Total cost of revenue",
     lambda i, c: f"=SUM({c}{rows['c_review']}:{c}{rows['c_pay']})", bold=True, top=True)
line("gp", "Gross profit", lambda i, c: f"={R('rev', c)}-{R('cogs', c)}", bold=True)
line("gm", "Gross margin %", lambda i, c: f"=IF({R('rev', c)}=0,0,{R('gp', c)}/{R('rev', c)})", PCT)

head("Operating expenses")
line("payroll", "Payroll (incl. burden)",
     lambda i, c: f"=SUMPRODUCT(--({START}<={c}$3),{SAL})*(1+{ref['burden']})")
line("mkt", "Marketing", lambda i, c: f"={R('d_new', c)}*{ref['cac']}+{ref['mkt_fixed']}")
line("ga", "General & administrative", lambda i, c: f"={ref['ga_base']}+{R('clients', c)}*{ref['ga_client']}")
line("opex", "Total operating expenses",
     lambda i, c: f"=SUM({c}{rows['payroll']}:{c}{rows['ga']})", bold=True, top=True)
line("ebitda", "EBITDA (operating profit)", lambda i, c: f"={R('gp', c)}-{R('opex', c)}", bold=True)

head("Cash")
line("cash_open", "Opening cash", lambda i, c: "=0")  # rewritten below once the closing-cash row exists
r -= 1  # placeholder rewritten below once cash_close row is known
cash_open_row = r
r += 1
line("funding", "Funding received", lambda i, c: f"=IF({c}$3={ref['raise_m']},{ref['raise']},0)")
line("cash_close", "Closing cash", lambda i, c: f"={c}{cash_open_row}+{R('ebitda', c)}+{R('funding', c)}", bold=True, top=True)
for i in range(MONTHS):
    c = m.cell(row=cash_open_row, column=FIRST_COL + i)
    c.value = f"={ref['cash0']}" if i == 0 else f"={L(i - 1)}{rows['cash_close']}"
    c.font = GREEN if i == 0 else BLACK
    c.number_format = USD
rows["cash_open"] = cash_open_row
line("burn", "Net burn (negative EBITDA)", lambda i, c: f"=MAX(0,-{R('ebitda', c)})")

head("Helper flags (used by Summary)")
line("f_cashout", "Month cash goes negative (else 999)", lambda i, c: f"=IF({R('cash_close', c)}<0,{c}$3,999)", '0')
line("f_profit", "Month EBITDA turns positive (else 999)", lambda i, c: f"=IF({R('ebitda', c)}>0,{c}$3,999)", '0')

last = L(MONTHS - 1)
def rng(key):
    return f"Monthly!$C${rows[key]}:${last}${rows[key]}"

# ---------------------------------------------------------------- Summary
s = wb.create_sheet("Summary", 0)
s.column_dimensions["A"].width = 40
for col in "BCD":
    s.column_dimensions[col].width = 16
s.column_dimensions["E"].width = 50
s["A1"] = "Closebook — 3-Year Summary ($)"
s["A1"].font = TITLE
s["A2"] = "Every number on this sheet is a formula driven by the Assumptions sheet."
s["A2"].font = Font(name="Arial", size=9, italic=True)

hdr = ["Metric", "Year 1", "Year 2", "Year 3"]
for j, h in enumerate(hdr):
    c = s.cell(row=4, column=1 + j, value=h)
    c.font = HEAD
    c.fill = HEAD_FILL
YR = f"Monthly!$C$4:${last}$4"

srow = 5
def srow_flow(label, key, fmt=USD, bold=False):
    global srow
    s.cell(row=srow, column=1, value=label).font = BOLD if bold else BLACK
    for y in range(3):
        c = s.cell(row=srow, column=2 + y, value=f"=SUMIF({YR},{y + 1},{rng(key)})")
        c.number_format = fmt
        c.font = BOLD if bold else BLACK
    srow += 1
    return srow - 1

def srow_point(label, key, fmt=NUM, bold=False):
    global srow
    s.cell(row=srow, column=1, value=label).font = BOLD if bold else BLACK
    for y in range(3):
        c = s.cell(row=srow, column=2 + y, value=f"=INDEX({rng(key)},{12 * (y + 1)})")
        c.number_format = fmt
        c.font = BOLD if bold else BLACK
    srow += 1
    return srow - 1

def srow_formula(label, fn, fmt, bold=False):
    global srow
    s.cell(row=srow, column=1, value=label).font = BOLD if bold else BLACK
    for y in range(3):
        col = get_column_letter(2 + y)
        c = s.cell(row=srow, column=2 + y, value=fn(col))
        c.number_format = fmt
        c.font = BOLD if bold else BLACK
    srow += 1
    return srow - 1

r_clients = srow_point("Clients at year end", "clients", NUM, True)
r_partners = srow_point("CPA partner firms at year end", "partners", NUM)
r_newc = srow_flow("New clients won", "new", NUM)
r_mrr = srow_point("Exit MRR (subscription)", "subs", USD)
srow_formula("Exit ARR (MRR × 12)", lambda col: f"={col}{r_mrr}*12", USD, True)
srow += 1
r_rev = srow_flow("Net revenue", "rev", USD, True)
r_gp = srow_flow("Gross profit", "gp", USD)
srow_formula("Gross margin %", lambda col: f"=IF({col}{r_rev}=0,0,{col}{r_gp}/{col}{r_rev})", PCT)
srow_flow("Operating expenses", "opex", USD)
srow_flow("EBITDA", "ebitda", USD, True)
srow_point("Closing cash", "cash_close", USD, True)
srow += 1

s.cell(row=srow, column=1, value="Key results").font = BOLD
for col in range(1, 5):
    s.cell(row=srow, column=col).fill = SECTION_FILL
srow += 1
def kv(label, formula, fmt, note=""):
    global srow
    s.cell(row=srow, column=1, value=label).font = BLACK
    c = s.cell(row=srow, column=2, value=formula)
    c.number_format = fmt
    c.font = BOLD
    s.cell(row=srow, column=5, value=note).font = Font(name="Arial", size=9, color="555555")
    srow += 1

kv("Lowest cash balance", f"=MIN({rng('cash_close')})", USD, "If negative, the raise is too small or costs too high.")
kv("Month cash runs out", f'=IF(MIN({rng("f_cashout")})=999,"Not within 36 months",MIN({rng("f_cashout")}))', '0', "")
kv("First month with positive EBITDA", f'=IF(MIN({rng("f_profit")})=999,"Not within 36 months",MIN({rng("f_profit")}))', '0', "")
kv("Peak monthly burn", f"=MAX({rng('burn')})", USD, "")
kv("Blended price per client (ARPA)", f"={ref['arpa']}", USD, "")
s.cell(row=srow - 1, column=2).font = GREEN

# ---------------------------------------------------------------- Unit economics
u = wb.create_sheet("Unit Economics", 1)
u.column_dimensions["A"].width = 46
u.column_dimensions["B"].width = 16
u.column_dimensions["C"].width = 56
u["A1"] = "Closebook — Unit Economics (per client, steady state)"
u["A1"].font = TITLE
urow = 3
uref = {}
def ue(key, label, formula, fmt, note="", link=False, bold=False):
    global urow
    u.cell(row=urow, column=1, value=label).font = BOLD if bold else BLACK
    c = u.cell(row=urow, column=2, value=formula)
    c.number_format = fmt
    c.font = GREEN if link else (BOLD if bold else BLACK)
    u.cell(row=urow, column=3, value=note).font = Font(name="Arial", size=9, color="555555")
    uref[key] = f"B{urow}"
    urow += 1

ue("arpa", "Revenue per client per month (ARPA)", f"={ref['arpa']}", USD, link=True)
ue("fee", "Payment processing per client", f"={uref['arpa']}*{ref['pay_fee']}", USD)
ue("cost", "Service cost per client (CPA + AI + data)", f"={ref['c_cpa']}+{ref['c_ai']}+{ref['c_data']}", USD)
ue("contrib", "Contribution per client per month", f"={uref['arpa']}-{uref['fee']}-{uref['cost']}", USD, bold=True)
ue("cm", "Contribution margin %", f"=IF({uref['arpa']}=0,0,{uref['contrib']}/{uref['arpa']})", PCT, "Deck target gross margin: 85–90%.")
ue("life", "Expected client lifetime (months)", f"=IF({ref['churn']}=0,0,1/{ref['churn']})", NUM1, "1 ÷ monthly churn.")
ue("ltv", "Lifetime value (contribution)", f"={uref['contrib']}*{uref['life']}", USD, "Excludes catch-up revenue, so it is conservative.", bold=True)
ue("cac", "Direct CAC", f"={ref['cac']}", USD, link=True)
ue("ratio", "LTV ÷ CAC", f"=IF({uref['cac']}=0,0,{uref['ltv']}/{uref['cac']})", '0.0"x"', "Above 3x is generally considered healthy.", bold=True)
ue("payback", "CAC payback (months)", f"=IF({uref['contrib']}<=0,0,{uref['cac']}/{uref['contrib']})", NUM1, "Months of contribution to earn back acquisition cost.")

for ws in (s, u, a, m):
    ws.sheet_view.showGridLines = False
wb.save(OUT)
print("saved", OUT)
